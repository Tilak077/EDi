"""Build the IEEE-style paper draft as an editable Word document.

    ./.venv/Scripts/python.exe build_paper_docx.py

Imports build_paper_pdf and walks the very same TITLE_BLOCK and STORY, so the
Word draft cannot drift away from the PDF: same paragraphs, same four tables,
same two figures, same numbers read from backend/experiments/results.json.

Layout: the title block and abstract run full width, then a continuous
two-column section (0.72 in margins, 0.24 in gutter -> 3.41 in columns)
carries the body. That is the same geometry as the PDF, so the figures and
tables, which were sized for one column, sit correctly without resizing.
"""

from __future__ import annotations

import html
import io
import re
from pathlib import Path

import build_paper_pdf as bp  # noqa: F401  (builds STORY, TITLE_BLOCK, REFS)

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

HERE = Path(__file__).parent
OUT = HERE / "MindWeave-R_IEEE_Paper_Draft.docx"
FIGDIR = HERE / "tmp" / "paper_docx"

TOKEN = re.compile(r"(<[^>]+>)")
COL_W = 3.41  # one column in inches; matches CW in build_paper_pdf.py
GUTTER_TWIPS = int(0.24 * 1440)

# reportlab ParagraphStyle name -> Word paragraph style name
STYLE_MAP = {
    "title": "MW Title",
    "author": "MW Author",
    "affil": "MW Affiliation",
    "abstract": "MW Abstract",
    "H1": "Heading 1",
    "H2": "Heading 2",
    "B": "MW Body",
    "B1": "MW Body Indent",
    "EQ": "MW Equation",
    "CAP": "MW Caption",
    "TCAP": "MW Table Label",
    "REF": "MW Reference",
}

COUNTS = {"figures": 0, "tables": 0, "paragraphs": 0}


# ---------------------------------------------------------------------- fonts
def force_font(style, name: str = "Times New Roman") -> None:
    """Pin every script of the style to the font, so Greek symbols and dashes
    do not fall back to a theme font."""
    rPr = style.element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.insert(0, rFonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rFonts.set(qn(attr), name)


def make_styles(doc: Document) -> None:
    def add(name: str, *, size: float, bold=False, italic=False, align=None,
            before=0.0, after=5.0, left=None, hang=None, keep=False,
            base="Normal"):
        st = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        st.base_style = doc.styles[base]
        st.font.name = "Times New Roman"
        st.font.size = Pt(size)
        if bold:
            st.font.bold = True
        if italic:
            st.font.italic = True
        st.font.color.rgb = RGBColor(0, 0, 0)
        pf = st.paragraph_format
        if align is not None:
            pf.alignment = align
        pf.space_before = Pt(before)
        pf.space_after = Pt(after)
        pf.keep_with_next = keep
        if left is not None:
            pf.left_indent = Inches(left)
        if hang is not None:
            pf.first_line_indent = Inches(hang)
        force_font(st)
        return st

    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(10)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.space_after = Pt(5)
    force_font(normal)

    add("MW Title", size=15, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=6)
    add("MW Author", size=11, align=WD_ALIGN_PARAGRAPH.CENTER, after=1)
    add("MW Affiliation", size=9.5, italic=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=1)
    add("MW Abstract", size=10, align=WD_ALIGN_PARAGRAPH.JUSTIFY, after=6)
    add("MW Body", size=10, align=WD_ALIGN_PARAGRAPH.JUSTIFY, after=5)
    add("MW Body Indent", size=10, align=WD_ALIGN_PARAGRAPH.JUSTIFY, after=4, left=0.2)
    add("MW Equation", size=10, align=WD_ALIGN_PARAGRAPH.CENTER, before=3, after=6)
    add("MW Caption", size=8.5, align=WD_ALIGN_PARAGRAPH.JUSTIFY, after=8)
    add("MW Table Label", size=8.5, align=WD_ALIGN_PARAGRAPH.CENTER, after=2, keep=True)
    add("MW Reference", size=9, align=WD_ALIGN_PARAGRAPH.JUSTIFY, after=2, left=0.2, hang=-0.2)
    add("MW Cell", size=8, align=WD_ALIGN_PARAGRAPH.CENTER, after=0)
    add("MW Cell Left", size=8, align=WD_ALIGN_PARAGRAPH.LEFT, after=0)

    for style_name, bold, italic, align in (
        ("Heading 1", True, False, WD_ALIGN_PARAGRAPH.CENTER),
        ("Heading 2", True, True, WD_ALIGN_PARAGRAPH.LEFT),
    ):
        st = doc.styles[style_name]
        st.font.name = "Times New Roman"
        st.font.size = Pt(10.5)
        st.font.bold = bold
        st.font.italic = italic
        st.font.color.rgb = RGBColor(0, 0, 0)
        pf = st.paragraph_format
        pf.alignment = align
        pf.space_before = Pt(10)
        pf.space_after = Pt(4)
        pf.keep_with_next = True
        force_font(st)


# -------------------------------------------------------------------- markup
def add_markup(par, text: str) -> None:
    """Write reportlab inline markup (<b>, <i>, <br/>, &mdash;) as Word runs."""
    bold = italic = False
    stack: list[tuple[bool, bool]] = []
    for tok in TOKEN.split(text):
        if not tok:
            continue
        if tok.startswith("<"):
            tag = tok.lower().replace(" ", "")
            if tag in ("<b>", "<strong>", "<i>", "<em>"):
                stack.append((bold, italic))
                if tag in ("<b>", "<strong>"):
                    bold = True
                else:
                    italic = True
            elif tag in ("</b>", "</strong>", "</i>", "</em>"):
                bold, italic = stack.pop() if stack else (False, False)
            elif tag in ("<br/>", "<br>"):
                par.add_run().add_break(WD_BREAK.LINE)
            else:
                raise ValueError(f"unsupported markup for Word export: {tok!r}")
        else:
            run = par.add_run(html.unescape(tok))
            if bold:
                run.bold = True
            if italic:
                run.italic = True


# ------------------------------------------------------------------ figures
def figure_png(flowable, path: Path, zoom: float = 6.0) -> Path:
    """Render one reportlab figure flowable to a PNG for the Word file."""
    import pymupdf
    from reportlab.pdfgen import canvas as rl_canvas

    w, h = flowable.width, flowable.height
    buf = io.BytesIO()
    c = rl_canvas.Canvas(buf, pagesize=(w + 4, h + 4))
    flowable.canv = c
    flowable.wrapOn(c, w + 4, h + 4)
    flowable.drawOn(c, 2, 2)
    c.showPage()
    c.save()

    pdf = pymupdf.open(stream=buf.getvalue(), filetype="pdf")
    pix = pdf[0].get_pixmap(matrix=pymupdf.Matrix(zoom, zoom), alpha=False)
    path.parent.mkdir(parents=True, exist_ok=True)
    pix.save(str(path))
    pdf.close()
    return path


def emit_figure(doc: Document, flowable, name: str) -> None:
    idx = 1 if name == "FigArchitecture" else 2
    png = figure_png(flowable, FIGDIR / f"fig{idx}.png")
    par = doc.add_paragraph(style="MW Body")
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    par.paragraph_format.keep_with_next = True
    par.paragraph_format.space_before = Pt(4)
    par.paragraph_format.space_after = Pt(2)
    par.add_run().add_picture(str(png), width=Inches(COL_W))
    COUNTS["figures"] += 1


# -------------------------------------------------------------------- tables
def cell_style_name(rl_style: str) -> str:
    return "MW Cell Left" if rl_style == "TDL" else "MW Cell"


def rule(tbl, edge: str, size: str, val: str) -> OxmlElement:
    el = OxmlElement(f"w:{edge}")
    el.set(qn("w:val"), val)
    el.set(qn("w:sz"), size)
    el.set(qn("w:space"), "0")
    el.set(qn("w:color"), "000000")
    return el


def emit_table(doc: Document, rl_table) -> None:
    data = getattr(rl_table, "_cellvalues", None)
    if data is None:
        data = getattr(rl_table, "cellvalues", None)
    if not data:
        raise ValueError("reportlab Table exposes no cell data")
    col_w = list(getattr(rl_table, "_colWidths", None) or [])
    n_cols = max(len(r) for r in data)
    if len(col_w) != n_cols:
        col_w = [1.0] * n_cols
    scale = COL_W * 72.0 / sum(col_w)  # fill the column exactly, keep proportions

    table = doc.add_table(rows=len(data), cols=n_cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    tbl = table._tbl
    tblPr = tbl.tblPr
    look = tblPr.find(qn("w:tblLook"))
    borders = OxmlElement("w:tblBorders")
    for edge, size, val in (
        ("top", "8", "single"), ("bottom", "8", "single"),
        ("insideH", "0", "none"), ("insideV", "0", "none"),
        ("left", "0", "none"), ("right", "0", "none"),
    ):
        borders.append(rule(tbl, edge, size, val))
    if look is not None:
        look.addprevious(borders)
    else:
        tblPr.append(borders)

    mar = OxmlElement("w:tblCellMar")
    for side, twips in (("top", "20"), ("left", "58"), ("bottom", "20"), ("right", "58")):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:w"), twips)
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    if look is not None:
        look.addprevious(mar)
    else:
        tblPr.append(mar)

    for r, row in enumerate(data):
        tr = table.rows[r]
        trPr = tr._tr.get_or_add_trPr()
        cant = OxmlElement("w:cantSplit")
        trPr.append(cant)
        if r == 0:
            trPr.append(OxmlElement("w:tblHeader"))
        for c in range(n_cols):
            cell = tr.cells[c]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell.width = Inches(col_w[c] * scale / 72.0)
            src = row[c] if c < len(row) else None
            par = cell.paragraphs[0]
            if src is None:
                continue
            style_name = src.style.name if getattr(src, "style", None) is not None else "TD"
            par.style = doc.styles[cell_style_name(style_name)]
            add_markup(par, src.text)
            if style_name == "TH":
                for run in par.runs:
                    run.bold = True
                # booktabs: the only interior rule sits under the header row
                tcPr = cell._tc.get_or_add_tcPr()
                tcBdr = OxmlElement("w:tcBorders")
                tcBdr.append(rule(None, "bottom", "4", "single"))
                tcPr.append(tcBdr)
    COUNTS["tables"] += 1

    gap = doc.add_paragraph()
    gap.paragraph_format.space_before = Pt(0)
    gap.paragraph_format.space_after = Pt(4)
    for run in gap.runs:
        run.font.size = Pt(4)


# ------------------------------------------------------------------ flowables
def emit_paragraph(doc: Document, rl_par) -> None:
    name = rl_par.style.name if rl_par.style is not None else None
    if name not in STYLE_MAP:
        raise ValueError(f"unmapped reportlab style for Word export: {name!r}")
    text = getattr(rl_par, "text", None)
    if not text:
        raise ValueError("reportlab Paragraph carries no source text")
    par = doc.add_paragraph(style=STYLE_MAP[name])
    add_markup(par, text)
    COUNTS["paragraphs"] += 1


def emit(doc: Document, flowables) -> None:
    for f in flowables:
        kind = type(f).__name__
        if isinstance(f, bp.KeepTogether):
            content = getattr(f, "_content", None) or getattr(f, "content", None)
            if content is None:
                raise ValueError("KeepTogether exposes no content")
            emit(doc, content)
        elif isinstance(f, bp.Paragraph):
            emit_paragraph(doc, f)
        elif isinstance(f, bp.Table):
            emit_table(doc, f)
        elif isinstance(f, bp.Spacer):
            continue
        elif kind in ("FigArchitecture", "FigRepairLoop"):
            emit_figure(doc, f, kind)
        elif kind == "NextPageTemplate":
            continue
        else:
            raise TypeError(f"unhandled flowable for Word export: {kind}")


# ------------------------------------------------------------------- sections
def configure_page(section) -> None:
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.left_margin = Inches(0.72)
    section.right_margin = Inches(0.72)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.62)
    section.footer_distance = Inches(0.35)
    section.gutter = Inches(0)


def set_columns(section, num: int, space_twips: int) -> None:
    sectPr = section._sectPr
    cols = sectPr.find(qn("w:cols"))
    if cols is None:
        cols = OxmlElement("w:cols")
        placed = False
        for succ in ("w:docGrid", "w:printerSettings", "w:sectPrChange"):
            el = sectPr.find(qn(succ))
            if el is not None:
                el.addprevious(cols)
                placed = True
                break
        if not placed:
            sectPr.append(cols)
    cols.set(qn("w:num"), str(num))
    cols.set(qn("w:space"), str(space_twips))
    cols.set(qn("w:equalWidth"), "1")


def add_footer(section) -> None:
    footer = section.footer
    footer.is_linked_to_previous = False
    paragraphs = footer.paragraphs
    par = paragraphs[0] if paragraphs else footer.add_paragraph()
    pPr = par._p.get_or_add_pPr()
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "center")
    tab.set(qn("w:pos"), str(int(7.06 * 1440 / 2)))
    tabs.append(tab)
    pPr.append(tabs)

    note = par.add_run("Working draft - not for distribution")
    note.font.size = Pt(8)
    note.font.color.rgb = RGBColor(0x59, 0x59, 0x59)
    par.add_run().add_tab()

    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    r = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")
    rFonts = OxmlElement("w:rFonts")
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rFonts.set(qn(attr), "Times New Roman")
    rPr.append(rFonts)
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), "16")
    rPr.append(sz)
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "595959")
    rPr.append(color)
    r.append(rPr)
    t = OxmlElement("w:t")
    t.text = "1"
    r.append(t)
    fld.append(r)
    par._p.append(fld)


# ----------------------------------------------------------------------- main
def main() -> None:
    doc = Document()
    configure_page(doc.sections[0])
    make_styles(doc)

    emit(doc, bp.TITLE_BLOCK)          # full-width title block and abstract
    sec = doc.add_section(WD_SECTION.CONTINUOUS)
    configure_page(sec)
    set_columns(sec, 2, GUTTER_TWIPS)  # body runs in two columns
    emit(doc, bp.STORY)

    add_footer(doc.sections[0])

    cp = doc.core_properties
    cp.title = "MindWeave-R: Contract-Aware Detection and Budgeted Repair of Artifacts in Multi-Agent Workflows"
    cp.author = "M. Barhate et al."
    cp.subject = "Working draft"

    doc.save(OUT)
    print(
        f"wrote {OUT} ({COUNTS['paragraphs']} paragraphs, {COUNTS['tables']} tables, "
        f"{COUNTS['figures']} figures, {len(bp.REF_NUM)} references cited)"
    )


if __name__ == "__main__":
    main()
